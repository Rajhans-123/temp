"""Builds the print-ready A4 Word report.

    python report/make_docx.py        -> report/dwm_report.docx

A4 (210 x 297 mm) portrait, one explicit page per section, so printing needs no
scaling: File -> Print -> Paper size A4, Margins: default, Scale: 100%. Numbers come
from report/content.py, which reads reports/*.csv - the same source the HTML preview
uses. Dashed frames are screenshot slots: paste a captured UI view over each one.
"""
from __future__ import annotations

import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

import content as K

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "dwm_report.docx")
FIG_DIR = os.path.join(HERE, "figures")

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
GREEN = RGBColor(0x16, 0x7A, 0x3A)
AMBER = RGBColor(0xB4, 0x53, 0x09)
BODY_FONT = "Segoe UI"
CODE_FONT = "Consolas"

CALLOUT = {
    "": (INDIGO, RGBColor(0xEE, 0xF2, 0xFF)),
    "ok": (GREEN, RGBColor(0xF0, 0xFD, 0xF4)),
    "warn": (AMBER, RGBColor(0xFF, 0xFB, 0xEB)),
}


# --------------------------------------------------------------- xml utils --
def shade(cell, hex_fill: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def cell_borders(cell, **edges) -> None:
    """edges: left=('single'|'dashed', size_eighths, hex), top=..., ..."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, spec in edges.items():
        if spec is None:
            continue
        style, size, colour = spec
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), colour)
        borders.append(el)
    tc_pr.append(borders)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def row_height(row, cm: float, exact: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"), str(int(cm * 567)))   # twips
    el.set(qn("w:hRule"), "exact" if exact else "atLeast")
    tr_pr.append(el)


def page_number_footer(section) -> None:
    footer = section.footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    for text, field in (("", "PAGE"), (" of ", None), ("", "NUMPAGES")):
        if field is None:
            run.add_text(text)
            continue
        run.font.name = BODY_FONT
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), field)
        inner = OxmlElement("w:r")
        fld.append(inner)
        run._r.addnext(fld)
    para.add_run("").font.size = Pt(1)


# ---------------------------------------------------------- inline markup --
_TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def add_rich(par, text: str, size: float = 9.5, colour: RGBColor = INK,
             bold: bool = False, italic: bool = False) -> None:
    """Renders `code`, **bold** and *italic* into runs."""
    for piece in _TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = par.add_run(piece[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(size - 0.8)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            run.bold = bold
        elif piece.startswith("**") and piece.endswith("**"):
            run = par.add_run(piece[2:-2])
            run.font.name = BODY_FONT
            run.font.size = Pt(size)
            run.font.color.rgb = colour
            run.bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run = par.add_run(piece[1:-1])
            run.font.name = BODY_FONT
            run.font.size = Pt(size)
            run.font.color.rgb = colour
            run.italic = True
        else:
            run = par.add_run(piece)
            run.font.name = BODY_FONT
            run.font.size = Pt(size)
            run.font.color.rgb = colour
            run.bold = bold
            run.italic = italic


def para(doc, text: str = "", size: float = 8.8, space_after: float = 3,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, colour: RGBColor = INK,
         italic: bool = False, bold: bool = False, space_before: float = 0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.0
    if text:
        add_rich(p, text, size=size, colour=colour, italic=italic, bold=bold)
    return p


def heading(doc, text: str, level: int):
    sizes = {1: 14.5, 2: 11, 3: 9.8}
    space_before = {1: 0, 2: 5, 3: 4}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before[level])
    p.paragraph_format.space_after = Pt(2 if level > 1 else 5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(sizes[level])
    run.font.color.rgb = INDIGO if level < 3 else INK
    run.bold = True
    if level == 1:
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), "4F46E5")
        pbdr.append(bottom)
        p._p.get_or_add_pPr().append(pbdr)
    return p


def listify(doc, items: list[str], ordered: bool = False, size: float = 8.7):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.65)
        pf.first_line_indent = Cm(-0.45)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.0
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(f"{i}." if ordered else "•")
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.font.color.rgb = INDIGO if ordered else INK
        run.bold = True
        add_rich(p, item, size=size)


def callout(doc, kind: str, text: str):
    accent, fill = CALLOUT.get(kind, CALLOUT[""])
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, f"{fill[0]:02X}{fill[1]:02X}{fill[2]:02X}"
                if isinstance(fill, tuple) else str(fill))
    cell_borders(cell, left=("single", 24, f"{accent[0]:02X}{accent[1]:02X}{accent[2]:02X}"),
                 top=("nil", 0, "auto"), bottom=("nil", 0, "auto"), right=("nil", 0, "auto"))
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.04
    add_rich(p, text, size=8.2)
    spacer(doc, 2)


def spacer(doc, points: float = 4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("")
    run.font.size = Pt(points)
    return p


def code_block(doc, text: str):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, "F8FAFC")
    cell_borders(cell, left=("single", 18, "4F46E5"), top=("single", 4, "E2E8F0"),
                 bottom=("single", 4, "E2E8F0"), right=("single", 4, "E2E8F0"))
    first = True
    for line in text.split("\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line)
        run.font.name = CODE_FONT
        run.font.size = Pt(7.8)


def data_table(doc, headers: list[str], rows: list, opts: dict):
    numeric = opts.get("numeric", set())
    code_cols = opts.get("code_cols", set())
    best = opts.get("best")
    size = opts.get("font", 7.8)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, "EEF2FF")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(head)
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.bold = True
        run.font.color.rgb = RGBColor(0x31, 0x2E, 0x81)
        if i in numeric:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_repeat_header(table.rows[0])

    for r, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if best is not None and r == best:
                shade(cell, "DCFCE7")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i in numeric:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            text = str(value)
            if i in code_cols:
                run = p.add_run(text)
                run.font.name = CODE_FONT
                run.font.size = Pt(size - 0.4)
            else:
                add_rich(p, text, size=size, bold=(best is not None and r == best))
    return table


def figure(doc, key: str, caption: str, width_cm: float = 13.6):
    path = os.path.join(FIG_DIR, f"fig_{key}.png")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    if os.path.exists(path):
        p.add_run().add_picture(path, width=Cm(width_cm))
    else:
        add_rich(p, f"[missing figure: {key}.png - run report/render_figures.py]",
                 size=8.5, colour=AMBER, italic=True)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(3)
    add_rich(cap, caption, size=7.4, colour=MUTED)


def shot_slots(doc, slots: list[tuple[str, str, float]]):
    """Pasteable frames: click the frame, Insert -> Pictures, then delete this text."""
    cols = len(slots)
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, (name, how, height_mm) in enumerate(slots):
        cell = table.cell(0, i)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, "F8FAFC")
        cell_borders(cell,
                     top=("dashed", 8, "94A3B8"), bottom=("dashed", 8, "94A3B8"),
                     left=("dashed", 8, "94A3B8"), right=("dashed", 8, "94A3B8"))
        tag = cell.paragraphs[0]
        tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tag.paragraph_format.space_after = Pt(1)
        run = tag.add_run("SCREENSHOT SLOT")
        run.font.name = BODY_FONT
        run.font.size = Pt(6.5)
        run.font.color.rgb = MUTED
        name_p = cell.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_p.paragraph_format.space_after = Pt(1)
        run = name_p.add_run(name)
        run.font.name = BODY_FONT
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        how_p = cell.add_paragraph()
        how_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        how_p.paragraph_format.space_after = Pt(0)
        run = how_p.add_run(how)
        run.font.name = BODY_FONT
        run.font.size = Pt(7.4)
        run.font.color.rgb = MUTED
        row_height(table.rows[0], height_mm / 10, exact=False)


# --------------------------------------------------------------- the cover --
def cover(doc, spec: dict):
    spacer(doc, 26)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(spec["course"].upper())
    run.font.name = BODY_FONT
    run.font.size = Pt(9.5)
    run.bold = True
    run.font.color.rgb = INDIGO

    rule = doc.add_table(rows=1, cols=1)
    cell = rule.cell(0, 0)
    shade(cell, "4F46E5")
    cell_borders(cell, top=("nil", 0, "auto"), bottom=("nil", 0, "auto"),
                 left=("nil", 0, "auto"), right=("nil", 0, "auto"))
    row_height(rule.rows[0], 0.12, exact=True)
    rule.rows[0].cells[0].width = Cm(4.5)
    spacer(doc, 14)

    for i, line in enumerate(spec["title"].split("\n")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2 if i == 0 else 8)
        run = p.add_run(line)
        run.font.name = BODY_FONT
        run.font.size = Pt(24)
        run.bold = True
        run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(spec["subtitle"])
    run.font.name = BODY_FONT
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.left_indent = Cm(1.6)
    p.paragraph_format.right_indent = Cm(1.6)
    p.paragraph_format.line_spacing = 1.15
    add_rich(p, spec["blurb"], size=9.8)

    kpi = doc.add_table(rows=2, cols=len(spec["kpis"]))
    kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (value, label, sub) in enumerate(spec["kpis"]):
        for r, (text, size, colour, bold) in enumerate(
                ((value, 13, INDIGO, True), (f"{label}\n{sub}", 7, MUTED, False))):
            cell = kpi.cell(r, i)
            shade(cell, "F8FAFC")
            cell_borders(cell, top=("single", 4, "E2E8F0"),
                         bottom=("single", 4, "E2E8F0"),
                         left=("single", 4, "E2E8F0"), right=("single", 4, "E2E8F0"))
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(0)
            run = cp.add_run(text)
            run.font.name = BODY_FONT
            run.font.size = Pt(size if r == 0 else 6.8)
            run.font.color.rgb = colour
            run.bold = bold
    spacer(doc, 16)

    meta = doc.add_table(rows=len(spec["rows"]), cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (key, value) in enumerate(spec["rows"]):
        left, right = meta.rows[i].cells
        for cell in (left, right):
            cell_borders(cell, top=("nil", 0, "auto"), bottom=("nil", 0, "auto"),
                         left=("nil", 0, "auto"), right=("nil", 0, "auto"))
        lp = left.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lp.paragraph_format.space_after = Pt(1.5)
        run = lp.add_run(key)
        run.font.name = BODY_FONT
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = MUTED
        rp = right.paragraphs[0]
        rp.paragraph_format.space_after = Pt(1.5)
        add_rich(rp, value, size=9)
    left.width = Cm(4.2)


def toc_table(doc, rows: list[tuple[str, str, str]]):
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for number, title, pages in rows:
        cells = table.add_row().cells
        for cell in cells:
            cell_borders(cell, top=("nil", 0, "auto"), bottom=("single", 4, "F1F5F9"),
                         left=("nil", 0, "auto"), right=("nil", 0, "auto"))
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(number)
        run.font.name = BODY_FONT
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = INDIGO
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_rich(p, title, size=9)
        p = cells[2].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(pages)
        run.font.name = BODY_FONT
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED


# ------------------------------------------------------------------ build --
def build() -> str:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(8.8)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(13)
    section.bottom_margin = Mm(12)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    page_number_footer(section)

    for page_no, blocks in enumerate(K.PAGES):
        # section starts use page_break_before on the heading itself: a separate
        # break paragraph can spill onto its own (blank) page when content is flush
        first = True
        for block in blocks:
            kind = block[0]
            started = None
            if kind == "cover":
                cover(doc, block[1])
            elif kind == "toc":
                toc_table(doc, block[1])
                spacer(doc, 6)
            elif kind in ("h1", "h2", "h3"):
                started = heading(doc, block[1], int(kind[1]))
            elif kind == "p":
                para(doc, block[1])
            elif kind == "pre":
                code_block(doc, block[1])
                spacer(doc, 4)
            elif kind == "callout":
                callout(doc, block[1], block[2])
            elif kind == "bullets":
                listify(doc, block[1], ordered=False)
            elif kind == "numbered":
                listify(doc, block[1], ordered=True)
            elif kind == "table":
                data_table(doc, block[1], block[2], block[3])
                spacer(doc, 5)
            elif kind == "figure":
                figure(doc, block[1], block[2],
                       width_cm=round(14.4 * block[3] / 100, 2))
            elif kind == "shots":
                shot_slots(doc, block[1])
            else:
                raise ValueError(f"unknown block {kind!r}")
            if first:
                first = False
                if page_no and started is not None:
                    started.paragraph_format.page_break_before = True

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    size = os.path.getsize(path) / 1024
    print(f"written: {path}  ({size:.0f} KB)")
    print("print: File -> Print -> Paper size A4, Margins default, Scale 100%")